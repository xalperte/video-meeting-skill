"""Renderer tests. Each is skipped if its library is unavailable. The PDF/ODP
conversions additionally need LibreOffice; we only assert on the always-present
outputs (.xlsx/.pptx/.docx) and treat converted files as best-effort."""
import importlib.util
import os
import subprocess
import sys
import tempfile
import unittest

import fixtures as F


def have(mod):
    return importlib.util.find_spec(mod) is not None


def run(script, args):
    return subprocess.run([sys.executable, os.path.join(F.SCRIPTS, script)] + args,
                          capture_output=True, text=True)


class TestRenderers(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()
        self.record = F.write_json(os.path.join(self.d, "record.json"),
                                   F.sample_record())

    @unittest.skipUnless(have("openpyxl"), "openpyxl not installed")
    def test_xlsx(self):
        out = os.path.join(self.d, "tasks.xlsx")
        p = run("render_tasks_xlsx.py", ["--record", self.record, "--out", out])
        self.assertEqual(p.returncode, 0, p.stderr)
        self.assertTrue(os.path.isfile(out))
        import openpyxl
        wb = openpyxl.load_workbook(out)
        self.assertIn("Action Items", wb.sheetnames)
        self.assertEqual(wb["Action Items"]["A2"].value, "Write the CSV exporter")

    @unittest.skipUnless(have("pptx"), "python-pptx not installed")
    def test_pptx(self):
        out = os.path.join(self.d, "slides.pptx")
        p = run("render_slides.py", ["--record", self.record, "--out-pptx", out,
                                     "--formats", "pptx"])
        self.assertEqual(p.returncode, 0, p.stderr)
        self.assertTrue(os.path.isfile(out))
        from pptx import Presentation
        self.assertGreaterEqual(len(Presentation(out).slides._sldIdLst), 3)

    @unittest.skipUnless(have("docx"), "python-docx not installed")
    def test_report_docx(self):
        out = os.path.join(self.d, "report.pdf")
        p = run("render_report.py", ["--record", self.record, "--out-pdf", out])
        self.assertEqual(p.returncode, 0, p.stderr)
        # the .docx is always produced; the .pdf only if LibreOffice is present
        self.assertTrue(os.path.isfile(os.path.join(self.d, "report.docx")))


class TestEmail(unittest.TestCase):
    def test_email_no_llm(self):
        with tempfile.TemporaryDirectory() as d:
            rec = F.write_json(os.path.join(d, "record.json"), F.sample_record())
            out = os.path.join(d, "email.md")
            p = run("render_email.py", ["--record", rec, "--out", out, "--no-llm"])
            self.assertEqual(p.returncode, 0, p.stderr)
            text = F.read_text(out)
            self.assertTrue(text.startswith("Subject:"))
            self.assertIn("Write the CSV exporter", text)


if __name__ == "__main__":
    unittest.main()
